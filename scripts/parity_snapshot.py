"""Half of the Python/browser parity harness: run every Python engine over a database and
write down exactly what it produced.

The tool is two implementations of the same logic — `src/` in Python and `docs/src/` in
JavaScript — kept identical by hand. That has already failed once in production: the
per-entity mapping fix shipped to Python and silently not to the browser app, which is the
copy the user actually works in. Nothing catches that class of drift, because the project has
no JavaScript test runner and no Node installed.

This is the cheapest thing that does catch it. Run this script, then open /parity.html in the
browser app: it loads the same database, runs the JavaScript engines, and diffs against the
snapshot written here. Two implementations, one database, every number compared.

    python scripts/parity_snapshot.py "path/to/consolidation.db"
    # then open docs/parity.html (serve docs/ and browse to /parity.html)

Formulas are compared as text rather than as evaluated numbers, because two workbooks can
agree on today's figures while referencing different cells.

The snapshot and the copied database land in docs/_parity/, which is gitignored — a real
consolidation database is client financial data and must not be committed.
"""

import json
import shutil
import sqlite3
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from src.db.repositories import consolidated_repo, group_coa_repo  # noqa: E402
from src.engines import comparative_engine, re_check               # noqa: E402
from src.engines.statement_generator import compute_bs, compute_pl # noqa: E402
from src.engines.validation_engine import run_all                  # noqa: E402
from src.exporters.excel_pack_generator import export_pack         # noqa: E402

OUT_DIR = PROJECT_ROOT / "docs" / "_parity"
PACK_SHEETS = ["PL_Current", "BS_Current", "Conso_TB_By_Period", "PL_Comparative", "BS_Comparative",
               "Conso_TB_Matrix", "Conso_TB_Total"]


def _round(value, places=4):
    return None if value is None else round(float(value), places)


def _statement(lines):
    return [{"label": l.label, "kind": l.kind, "level": l.level,
             "value": _round(l.value), "percent": _round(l.percent, 10)} for l in lines]


def _comparative(result):
    with_data = [c for c in result.columns if result.has_data(c)]
    leaves = [l for l in result.lines if l.kind != "section"]
    return {
        "columns": result.columns,
        "withData": with_data,
        "missing": result.context.missing,
        "open": result.context.open_in_range,
        "stale": result.context.stale_in_range,
        "lines": [[l.label, l.kind, l.level] for l in result.lines],
        "values": {c: {l.label: _round(result.value(c, l.label)) for l in leaves} for c in with_data},
        "percents": {c: {l.label: _round(result.percent(c, l.label), 10) for l in leaves} for c in with_data},
        "totals": {c: _round(v) for c, v in result.totals_by_column.items()},
    }


def _buckets(conn, period_id):
    by_entity, group_adj, total = consolidated_repo.get_buckets(conn, period_id)
    return {
        "byAccountEntity": {f"{a}:{e}": _round(v) for (a, e), v in by_entity.items()},
        "groupAdjustment": {str(a): _round(v) for a, v in group_adj.items()},
        "total": {str(a): _round(v) for a, v in total.items()},
    }


def _re_check(conn, period_id):
    result = re_check.run(conn, period_id)
    return {
        "applicable": result.applicable,
        "reason": result.reason,
        "priorPeriodStr": result.prior_period_str,
        "rows": [{"entityCode": r.entity_code, "reClosingPrior": _round(r.re_closing_prior),
                  "priorNetProfit": _round(r.prior_net_profit),
                  "reClosingCurrent": _round(r.re_closing_current),
                  "expected": _round(r.expected), "gap": _round(r.gap)} for r in result.rows],
        "linkedAdjustmentCount": len(result.linked_adjustments),
    }


def _validations(conn, period_id):
    return [{"checkId": v.check_id, "description": v.description, "severity": v.severity,
             "passed": v.passed, "details": list(v.details)} for v in run_all(conn, period_id)]


def _matrix(conn, period_id):
    context, columns, totals = comparative_engine.consolidated_matrix(conn, period_id)
    by_code = {r["group_account_id"]: r["account_code"] for r in group_coa_repo.list_all(conn)}
    return {
        "columns": columns,
        "missing": context.missing,
        "open": context.open_in_range,
        "cells": {c: {by_code[a]: _round(v) for a, v in totals[c].items() if a in by_code}
                  for c in columns},
    }


def _pack_cells(conn, period_id, tmp_path):
    """Every cell of the statement sheets, formulas as text. openpyxl stores a formula as a
    string starting with '=' while ExcelJS stores {formula} without it, so the '=' is dropped
    here and the browser side compares the bare formula."""
    import openpyxl

    export_pack(conn, period_id, str(tmp_path))
    wb = openpyxl.load_workbook(tmp_path, data_only=False)
    out = {"sheetNames": wb.sheetnames, "cells": {}}
    for name in PACK_SHEETS:
        if name not in wb.sheetnames:
            continue
        cells = {}
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cells[cell.coordinate] = {"f": cell.value[1:]}
                elif isinstance(cell.value, (int, float)):
                    cells[cell.coordinate] = {"v": _round(cell.value)}
                else:
                    cells[cell.coordinate] = {"s": str(cell.value)}
        out["cells"][name] = cells
    return out


def build(db_path: str) -> dict:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    periods = [dict(r) for r in conn.execute(
        "SELECT period_id, year, month, status FROM periods ORDER BY year, month")]

    snapshot = {"periods": [], "generatedFrom": Path(db_path).name}
    for p in periods:
        pid = p["period_id"]
        label = f"{p['year']}-{p['month']:02d}"
        pl_lines, net_income = compute_pl(conn, pid)
        bs_lines, total_assets, total_le = compute_bs(conn, pid)
        entry = {
            "periodId": pid,
            "label": label,
            "pl": {"lines": _statement(pl_lines), "netIncome": _round(net_income)},
            "bs": {"lines": _statement(bs_lines), "totalAssets": _round(total_assets),
                   "totalLe": _round(total_le)},
            "buckets": _buckets(conn, pid),
            "plComparative": _comparative(comparative_engine.compute_pl_comparative(conn, pid)),
            "bsComparative": _comparative(comparative_engine.compute_bs_comparative(conn, pid)),
            "matrix": _matrix(conn, pid),
            "validations": _validations(conn, pid),
            "reCheck": _re_check(conn, pid),
        }
        snapshot["periods"].append(entry)

    # The Excel pack is slow to build, so only the newest period gets one. That is the period
    # a pack would actually be exported for.
    if periods:
        newest = periods[-1]["period_id"]
        snapshot["pack"] = {"periodId": newest,
                            **_pack_cells(conn, newest, OUT_DIR / "_pack.xlsx")}
        (OUT_DIR / "_pack.xlsx").unlink(missing_ok=True)

    conn.close()
    return snapshot


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    db_path = argv[1]
    if not Path(db_path).exists():
        print(f"No such database: {db_path}")
        return 2

    snapshot = build(db_path)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "expected.json").write_text(json.dumps(snapshot), encoding="utf-8")
    # The browser side needs the identical database, served from the same origin as the app.
    shutil.copy2(db_path, OUT_DIR / "fixture.db")

    periods = ", ".join(p["label"] for p in snapshot["periods"])
    print(f"Snapshot written to {OUT_DIR / 'expected.json'}")
    print(f"  periods: {periods or '(none)'}")
    if "pack" in snapshot:
        total = sum(len(v) for v in snapshot["pack"]["cells"].values())
        print(f"  pack cells: {total} across {len(snapshot['pack']['cells'])} sheets")
    print(f"  fixture database copied to {OUT_DIR / 'fixture.db'}")
    print("\nNow serve docs/ and open /parity.html to run the browser side.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
