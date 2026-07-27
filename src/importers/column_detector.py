"""Locates the Tally TB header block by searching for known marker text rather than
assuming a fixed row/column position (source plan 6.1 requirement: flexible column
detection since Tally exports may shift slightly)."""

from dataclasses import dataclass
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class ColumnLayout:
    particulars_col: int
    opening_col: int
    period_debit_col: int
    period_credit_col: int
    closing_col: int
    data_start_row: int


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def detect(ws: Worksheet) -> ColumnLayout:
    particulars_row = None
    particulars_col = None
    for row in range(1, min(ws.max_row, 30) + 1):
        for col in range(1, min(ws.max_column, 15) + 1):
            if _norm(ws.cell(row=row, column=col).value) == "particulars":
                particulars_row, particulars_col = row, col
                break
        if particulars_row:
            break
    if particulars_row is None:
        raise ValueError("Could not locate a 'Particulars' header cell in this workbook")

    opening_col = period_start_col = closing_col = None
    marker_row = particulars_row + 1
    for col in range(1, min(ws.max_column, 20) + 1):
        val = _norm(ws.cell(row=marker_row, column=col).value)
        if val == "opening":
            opening_col = col
        elif val == "transactions":
            period_start_col = col
        elif val == "closing":
            closing_col = col
    if not all([opening_col, period_start_col, closing_col]):
        raise ValueError(
            f"Could not locate Opening/Transactions/Closing markers on row {marker_row}"
        )

    debit_col = credit_col = None
    sub_marker_row = particulars_row + 2
    for col in range(period_start_col, closing_col):
        val = _norm(ws.cell(row=sub_marker_row, column=col).value)
        if val == "debit":
            debit_col = col
        elif val == "credit":
            credit_col = col
    if not all([debit_col, credit_col]):
        raise ValueError(
            f"Could not locate Debit/Credit markers on row {sub_marker_row}"
        )

    return ColumnLayout(
        particulars_col=particulars_col,
        opening_col=opening_col,
        period_debit_col=debit_col,
        period_credit_col=credit_col,
        closing_col=closing_col,
        data_start_row=particulars_row + 3,
    )
