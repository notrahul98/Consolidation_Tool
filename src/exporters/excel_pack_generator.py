"""Excel pack: Summary, Conso_TB_Matrix, Conso_TB_Total, Entity_<CODE> sheets, PL_Current,
BS_Current, Conso_TB_By_Period, PL_Comparative, BS_Comparative, Mapping, Validation.

Every derived cell is a live Excel formula, not a static value, so the workbook is
traceable directly in Excel (click a cell, see where it comes from):

    Entity_<CODE> sheets (base data, from the DB)
        -> Conso_TB_Matrix (entity columns reference Entity_<CODE>; Adjustments holds
           GROUP-LEVEL adjustments only; Total = SUM across the entity columns AND
           the Adjustments column)
            -> Conso_TB_Total (Adjustments/Total columns reference Conso_TB_Matrix's own)
                -> PL_Current / BS_Current (leaves reference Conso_TB_Total; subtotals are
                   SUM/arithmetic formulas over rows within the same sheet; BS_Current's
                   Retained Earnings cross-references PL_Current's Net Income row directly,
                   same as the real template's Schedule 15 pulls in the current year's P&L)
                -> Conso_TB_By_Period (one column per month; the anchor month's column is a
                   formula into Conso_TB_Total so the two can never disagree, other months are
                   written data, and the two YTD columns are SUMs over a month range)
                    -> PL_Comparative / BS_Comparative (same rows and the same subtotal
                       formulas as PL_Current / BS_Current, one column per period)

statement_generator.py's compute_pl/compute_bs (numeric, not formulas) remain the source of
truth for the validation engine — this module's formulas are built to mirror that logic
exactly and are checked against it via LibreOffice recalculation, not by sharing code, since
one writes Python values and the other writes Excel formula strings. tests/test_excel_pack.py
evaluates the whole chain and diffs it against compute_pl/compute_bs, which is what catches a
drift like the Total column silently excluding group-level adjustments.
"""

import sqlite3

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.db.repositories import entity_repo, group_coa_repo, consolidated_repo
from src.engines import comparative_engine
from src.engines.validation_engine import run_all
from src.engines.statement_generator import compute_pl, compute_bs

IDR_FORMAT = '#,##0;(#,##0)'
BOLD = Font(bold=True)
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")


def _write_summary(wb, conn, period_id, period_label, validations):
    ws = wb.create_sheet("Summary")
    ws.append(["Tally Consolidation Tool - Phase 1"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Period", period_label])
    ws.append([])
    ws.append(["Entities"])
    ws[f"A{ws.max_row}"].font = BOLD
    for e in entity_repo.list_all(conn):
        ws.append([e["entity_code"], e["entity_name"]])
    ws.append([])
    ws.append(["Validation status"])
    ws[f"A{ws.max_row}"].font = BOLD
    passed = sum(1 for v in validations if v.passed)
    ws.append([f"{passed}/{len(validations)} checks passed"])
    for v in validations:
        row = [v.check_id, v.description, "PASS" if v.passed else "FAIL"]
        ws.append(row)
        fill = GREEN_FILL if v.passed else RED_FILL
        ws.cell(row=ws.max_row, column=3).fill = fill
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 55


def _write_entity_sheets(wb, conn, period_id):
    """Base layer: every leaf category for every entity, in identical canonical order
    across all entity sheets (even zero-balance ones) so row numbers are predictable and
    Conso_TB_Matrix can reference them by cell, not by re-deriving the value."""
    entities = entity_repo.list_all(conn)
    coa_leaves = group_coa_repo.list_leaf_categories(conn)
    # Shared with the Consolidated TB page. Each entity's cell is its TB plus any adjustment
    # booked against that entity, summed (see HANDOVER.md §9.6). Group-level adjustments are
    # deliberately NOT here — they belong to no entity and are carried by the matrix's own
    # Adjustments column instead.
    by_account_entity, _, _ = consolidated_repo.get_buckets(conn, period_id)

    entity_row: dict[str, dict[str, int]] = {}
    for e in entities:
        ws = wb.create_sheet(f"Entity_{e['entity_code']}")
        ws.append(["Account Code", "Account Name", "Section", "Closing Balance"])
        for c in ws[1]:
            c.font = BOLD
        ws.freeze_panes = "A2"
        code_to_row: dict[str, int] = {}
        for row in coa_leaves:
            v = by_account_entity.get((row["group_account_id"], e["entity_id"]), 0.0)
            ws.append([row["account_code"], row["account_name"], row["section"], v])
            ws.cell(row=ws.max_row, column=4).number_format = IDR_FORMAT
            code_to_row[row["account_code"]] = ws.max_row
        entity_row[e["entity_code"]] = code_to_row
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 16

    return entity_row


def _write_matrix_and_total(wb, conn, period_id, coa_rows, entities, entity_row):
    """Conso_TB_Matrix: entity columns formula-reference Entity_<CODE> sheets; Adjustments is
    a static column holding GROUP-LEVEL adjustments only (there's no per-entity "Entity_ADJ"
    sheet to formula-reference, so this one column is a computed value rather than a
    cross-sheet formula). Total is a SUM formula across the entity columns AND the
    Adjustments column. Conso_TB_Total mirrors the same row layout and formula-references
    Conso_TB_Matrix's own Adjustments and Total columns.

    The Adjustments column must be group-level only, and the Total must span it:
      * entity-specific adjustments are already inside the Entity_<CODE> sheets, so counting
        them here as well would double them;
      * group-level adjustments live in no entity column at all, so leaving them outside the
        SUM drops them from the workbook entirely — the Balance Sheet then disagrees with the
        tool while Total Assets, Total Liabilities and Equity, and Check all still tie,
        because a balanced group-level journal hides inside the totals.
    """
    _, group_adjustment, _ = consolidated_repo.get_buckets(conn, period_id)

    ws = wb.create_sheet("Conso_TB_Matrix")
    header = ["Account Code", "Account Name", "Section"] + [e["entity_code"] for e in entities] \
        + ["Adjustments", "Total"]
    ws.append(header)
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "D2"

    entity_start_col = 4
    adj_col = entity_start_col + len(entities)
    total_col = adj_col + 1
    adj_col_letter = get_column_letter(adj_col)
    total_col_letter = get_column_letter(total_col)

    code_row: dict[str, int] = {}
    for row in coa_rows:
        ws.append([row["account_code"], row["account_name"], row["section"]])
        r = ws.max_row
        code_row[row["account_code"]] = r
        if not row["is_header"]:
            for i, e in enumerate(entities):
                col = entity_start_col + i
                entity_sheet_row = entity_row[e["entity_code"]].get(row["account_code"])
                cell = ws.cell(row=r, column=col)
                if entity_sheet_row:
                    cell.value = f"='Entity_{e['entity_code']}'!D{entity_sheet_row}"
                cell.number_format = IDR_FORMAT
            adj_cell = ws.cell(row=r, column=adj_col)
            adj_cell.value = group_adjustment.get(row["group_account_id"]) or None
            adj_cell.number_format = IDR_FORMAT

            # Spans the entity columns THROUGH the Adjustments column — see the note above.
            first_letter = get_column_letter(entity_start_col)
            ws.cell(row=r, column=total_col, value=f"=SUM({first_letter}{r}:{adj_col_letter}{r})")
            ws.cell(row=r, column=total_col).number_format = IDR_FORMAT
        else:
            for c in ws[r]:
                c.font = BOLD

    widths = [12, 40, 20] + [14] * (len(entities) + 2)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws_total = wb.create_sheet("Conso_TB_Total")
    ws_total.append(["Account Code", "Account Name", "Section", "Adjustments", "Total"])
    for c in ws_total[1]:
        c.font = BOLD
    ws_total.freeze_panes = "D2"
    for row in coa_rows:
        ws_total.append([row["account_code"], row["account_name"], row["section"]])
        r = ws_total.max_row
        assert r == code_row[row["account_code"]], "Conso_TB_Total must mirror Conso_TB_Matrix row-for-row"
        if not row["is_header"]:
            ws_total.cell(row=r, column=4, value=f"='Conso_TB_Matrix'!{adj_col_letter}{r}")
            ws_total.cell(row=r, column=4).number_format = IDR_FORMAT
            ws_total.cell(row=r, column=5, value=f"='Conso_TB_Matrix'!{total_col_letter}{r}")
            ws_total.cell(row=r, column=5).number_format = IDR_FORMAT
        else:
            for c in ws_total[r]:
                c.font = BOLD
    for i, w in enumerate([12, 40, 20, 16, 16], start=1):
        ws_total.column_dimensions[get_column_letter(i)].width = w

    return code_row


def _write_mapping_sheet(wb, conn, period_id):
    from src.db.repositories import mapping_repo
    ws = wb.create_sheet("Mapping")
    ws.append(["Ledger Name", "Entities Present In", "Category", "Status"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"
    for entry in sorted(mapping_repo.distinct_ledgers_for_period(conn, period_id), key=lambda x: x["ledger_name"]):
        category = ""
        if len(entry["group_account_ids"]) == 1:
            row = group_coa_repo.get_by_id(conn, next(iter(entry["group_account_ids"])))
            category = row["account_name"] if row else ""
        status = "Mapped" if entry["group_account_ids"] else "Unmapped"
        ws.append([entry["ledger_name"], ", ".join(sorted(set(entry["entities"]))), category, status])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 14


def _write_validation_sheet(wb, validations):
    ws = wb.create_sheet("Validation")
    ws.append(["Check", "Description", "Severity", "Status", "Details"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"
    for v in validations:
        detail_text = "; ".join(v.details) if v.details else ""
        ws.append([v.check_id, v.description, v.severity, "PASS" if v.passed else "FAIL", detail_text])
        fill = GREEN_FILL if v.passed else RED_FILL
        ws.cell(row=ws.max_row, column=4).fill = fill
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 80


def _leaf_ref(name, name_to_leaf, code_row, positive_override=False):
    """Formula fragment (no leading '=') pulling a leaf category from Conso_TB_Total,
    applying the Dr-positive -> display-magnitude sign flip via normal_balance."""
    info = name_to_leaf[name]
    row = code_row[info["account_code"]]
    negate = (info["normal_balance"] == "CR") and not positive_override
    ref = f"'Conso_TB_Total'!E{row}"
    return f"-{ref}" if negate else ref


# --- Statement layouts --------------------------------------------------------------------
# The P&L and BS row structures, written once and rendered into any column.
#
# PL_Current fills one column; PL_Comparative fills twenty-six. Defining the layout as data
# rather than as inline openpyxl calls is what keeps those two from being separate copies that
# drift — a subtotal formula changed in one and not the other would show up as a single wrong
# column in a fifty-column sheet, which is exactly the kind of thing nobody notices.
#
# Each row is a dict:
#   leaf    -> pull `name` from the trial balance, sign-flipped for display
#   group   -> header whose amount is SUM over its members' rows; `members_first` puts the
#              header after its members (the Balance Sheet's convention, unlike the P&L's)
#   calc    -> amount is fn(row_of, col, ctx) -> formula fragment
#   section -> a bare label with no amount

def _leaf(label, level=0, name=None, positive_override=False):
    return {"kind": "leaf", "label": label, "level": level, "name": name or label,
            "positive_override": positive_override, "bold": False}


def _calc(label, fn, bold=True, level=0):
    return {"kind": "calc", "label": label, "level": level, "fn": fn, "bold": bold}


def _group(label, members, members_first=False):
    return {"kind": "group", "label": label, "level": 0, "members": members,
            "members_first": members_first, "bold": True}


def _section(label):
    return {"kind": "section", "label": label, "level": 0, "bold": True}


def _expand(layout):
    """Flatten group entries into (header, members) in the right order for their statement."""
    rows = []
    for spec in layout:
        if spec["kind"] != "group":
            rows.append(spec)
            continue
        members = [_leaf(m, level=1) for m in spec["members"]]
        rows.extend(members + [spec] if spec["members_first"] else [spec] + members)
    return rows


def PL_LAYOUT():
    opex_members = ["Staff and Staff related Costs", "Marketing Expenses", "Sales Expenses", "Utlities",
                    "Profit/Loss Sharing", "Rent Expense", "Depreciation", "Consumables",
                    "Maintenance Expenses", "Taxes, Licenses & Permits", "Other Expenses", "Travel Expenses"]
    return _expand([
        _leaf("Sales", 1),
        _calc("Total Sales", lambda R, C, X: f"{C}{R['Sales']}"),
        _leaf("Sales Discount/Return", 1),
        _calc("Net Sales", lambda R, C, X: f"{C}{R['Total Sales']}-{C}{R['Sales Discount/Return']}"),

        _leaf("COGS before Direct Cost & Forex Gain/Loss", 1),
        _leaf("Direct Income", 1),
        _leaf("Direct Costs", 1),
        _calc("Total Cost of Goods Sold",
              lambda R, C, X: f"{C}{R['COGS before Direct Cost & Forex Gain/Loss']}"
                              f"+{C}{R['Direct Costs']}-{C}{R['Direct Income']}"),
        _calc("Gross Profit", lambda R, C, X: f"{C}{R['Net Sales']}-{C}{R['Total Cost of Goods Sold']}"),

        _group("Staff and Staff related Costs",
               ["Salaries & Wages", "Bonus & Incentives", "Commission", "Staff Insurance", "Staff Welfare"]),
        _group("Marketing Expenses", ["Advertising Expense", "Entertainment"]),
        _group("Sales Expenses", ["Freight, Storage & Handling", "Sales Fees", "Packaging and Labeling"]),
        _group("Utlities", ["Telephone Expense", "Electricity Expense"]),
        _leaf("Profit/Loss Sharing"),
        _leaf("Rent Expense"),
        _leaf("Depreciation"),
        _group("Consumables", ["Printing & Stationery", "Pantry Expenses"]),
        _group("Maintenance Expenses", ["Repairs & Maintenance", "Security Expense"]),
        _leaf("Taxes, Licenses & Permits"),
        _group("Other Expenses", [
            "Sample Expenses", "Interest Expense", "Interest on  Shareholder loan", "Warehouse Expense",
            "Consulting & Professional Fee", "General Insurance Expense", "Postage & Courier Expenses",
            "Tax Expenses", "Transport Expense", "Office Expenses", "Modern Market Trading Terms",
            "Stock Write Off", "Bad Debts",
        ]),
        _group("Travel Expenses", ["Travel Expense - International", "Travel Expense - Domestic"]),

        _calc("Operating Expenses", lambda R, C, X: "+".join(f"{C}{R[m]}" for m in opex_members)),
        _calc("Operative Profit/(Loss)", lambda R, C, X: f"{C}{R['Gross Profit']}-{C}{R['Operating Expenses']}"),

        _group("Miscellaneous Expenses", ["Misc. Expense", "Bank Charges", "Foreign Exchange (Gain) Loss"]),
        _group("Miscellaneous Incomes", ["Other Income", "Other Income - CPCI Income"]),
        _calc("Total Non Operating Expenses/(Income)",
              lambda R, C, X: f"{C}{R['Miscellaneous Expenses']}-{C}{R['Miscellaneous Incomes']}"),
        _calc("Net Income/(Loss) Before Tax & Previous Year Expenses",
              lambda R, C, X: f"{C}{R['Operative Profit/(Loss)']}-{C}{R['Total Non Operating Expenses/(Income)']}"),

        _group("Extraordinary / Previous Year Expenses", [
            "Tax Expenses - Prev Years", "Credit Note - Prev. Years", "Previous Year Rent Expenses",
            "Previous Year Commissions", "Corporate Tax",
        ]),
        _calc("Net Income/(Loss) After Tax & Previous Year Expenses",
              lambda R, C, X: f"{C}{R['Net Income/(Loss) Before Tax & Previous Year Expenses']}"
                              f"-{C}{R['Extraordinary / Previous Year Expenses']}"),

        _calc("Net Income Before Depreciation",
              lambda R, C, X: f"{C}{R['Net Income/(Loss) After Tax & Previous Year Expenses']}+{C}{R['Depreciation']}"),
        _calc("Net Profit without interest on bank loan",
              lambda R, C, X: f"{C}{R['Net Income/(Loss) After Tax & Previous Year Expenses']}+{C}{R['Interest Expense']}"),
        _calc("Net Profit without interest on shareholder loan",
              lambda R, C, X: f"{C}{R['Net Income/(Loss) After Tax & Previous Year Expenses']}"
                              f"+{C}{R['Interest on  Shareholder loan']}"),
    ])


def BS_LAYOUT():
    return _expand([
        _section("CURRENT ASSETS"),
        _group("Total Current Assets", [
            "Cash & Cash Equivalent", "Accounts Receivable", "Inventory", "Prepaid Taxes",
            "Prepaid Expenses", "Short Term Deposits", "Other Receivables",
        ], members_first=True),

        _section("NON CURRENT ASSETS"),
        _leaf("Fixed Assets", 1),
        # Displayed as the raw (already-negative) balance — a contra-asset, matching the real
        # template's convention — so Net Fixed Assets below is a plain addition, not a subtraction.
        _leaf("Accumulated Depreciation", 1, positive_override=True),
        _calc("Net Fixed Assets", lambda R, C, X: f"{C}{R['Fixed Assets']}+{C}{R['Accumulated Depreciation']}"),
        _leaf("Other Non-current Asset"),
        _calc("Total Non-Current Assets",
              lambda R, C, X: f"{C}{R['Net Fixed Assets']}+{C}{R['Other Non-current Asset']}"),
        _calc("TOTAL ASSETS", lambda R, C, X: f"{C}{R['Total Current Assets']}+{C}{R['Total Non-Current Assets']}"),

        _section("LIABILITIES AND EQUITY"),
        _section("LIABILITIES"),
        _group("Total Short Term Liabilities", [
            "Accounts Payable", "Taxes Payable", "Accrued Expenses", "Loans & Advances taken",
            "Interco Balances", "Other Payables",
        ], members_first=True),

        _section("EQUITY"),
        _leaf("Share Capital", 1),
        # Retained Earnings = brought-forward P&L ledger balance + this period's P&L result,
        # cross-referenced straight from the P&L sheet — same composition as the real template's
        # Schedule 15 (prior years + current year P&L). ctx supplies the reference because it
        # points at a different sheet and column for BS_Current than for BS_Comparative.
        _calc("Retained Earnings",
              lambda R, C, X: f"({X['re_leaf']})+{X['pl_net_income'](C)}", bold=False, level=1),
        _calc("Total Equity", lambda R, C, X: f"{C}{R['Share Capital']}+{C}{R['Retained Earnings']}"),
        _calc("TOTAL LIABILITIES AND EQUITY",
              lambda R, C, X: f"{C}{R['Total Short Term Liabilities']}+{C}{R['Total Equity']}"),
        _calc("Check", lambda R, C, X: f"{C}{R['TOTAL LIABILITIES AND EQUITY']}-{C}{R['TOTAL ASSETS']}"),
    ])


def _place_labels(ws, rows, label_col=1):
    """Write the row labels down column A at the sheet's current position; returns
    {label: sheet_row}. Rows are identical across every amount column, so this runs once."""
    row_of = {}
    for spec in rows:
        indent = "    " * spec["level"]
        ws.append([f"{indent}{spec['label']}"])
        r = ws.max_row
        if spec["bold"]:
            ws.cell(row=r, column=label_col).font = BOLD
        row_of[spec["label"]] = r
    return row_of


def _fill_column(ws, rows, row_of, col_letter, ref, ctx=None):
    """Write one amount column's formulas. `ref(name, positive_override)` returns the leaf
    reference fragment, which is the only thing that differs between a single-period sheet and
    one comparative column."""
    ctx = ctx or {}
    for spec in rows:
        if spec["kind"] == "section":
            continue
        r = row_of[spec["label"]]
        if spec["kind"] == "leaf":
            formula = ref(spec["name"], spec["positive_override"])
        elif spec["kind"] == "group":
            first = row_of[spec["members"][0]]
            last = row_of[spec["members"][-1]]
            formula = f"SUM({col_letter}{first}:{col_letter}{last})"
        else:
            formula = spec["fn"](row_of, col_letter, ctx)
        cell = ws[f"{col_letter}{r}"]
        cell.value = f"={formula}"
        cell.number_format = IDR_FORMAT
        if spec["bold"]:
            cell.font = BOLD


def _write_pl_sheet(wb, conn, name_to_leaf, code_row):
    ws = wb.create_sheet("PL_Current")
    ws.append(["Line Item", "Amount (IDR)", "% of Net Sales"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"

    rows = PL_LAYOUT()
    row_of = _place_labels(ws, rows)
    _fill_column(ws, rows, row_of, "B",
                 lambda name, pos: _leaf_ref(name, name_to_leaf, code_row, pos))

    net_sales_row = row_of["Net Sales"]
    for r in row_of.values():
        ws.cell(row=r, column=3, value=f"=B{r}/$B${net_sales_row}")
        ws.cell(row=r, column=3).number_format = "0.0%"

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 15

    return row_of["Net Income/(Loss) After Tax & Previous Year Expenses"]


def _write_bs_sheet(wb, conn, name_to_leaf, code_row, pl_net_income_row):
    ws = wb.create_sheet("BS_Current")
    ws.append(["Line Item", "Amount (IDR)"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"

    rows = BS_LAYOUT()
    row_of = _place_labels(ws, rows)
    ref = lambda name, pos: _leaf_ref(name, name_to_leaf, code_row, pos)
    _fill_column(ws, rows, row_of, "B", ref, ctx={
        "re_leaf": ref("Retained Earnings", False),
        "pl_net_income": lambda col: f"'PL_Current'!B{pl_net_income_row}",
    })

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 20


MONTH_ABBR = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _month_label(year, month):
    return f"{MONTH_ABBR[month]}'{str(year)[2:]}"


def _write_by_period_sheet(wb, conn, period_id, coa_leaves, code_row, context):
    """Conso_TB_By_Period: one column per month of the current and prior year, plus the two
    year-to-date columns, for every leaf category.

    This is the base layer the comparative statements reference, playing the same role for
    months that the Entity_<CODE> sheets play for entities. Months are laid out as two
    contiguous year blocks rather than interleaved pairs so each YTD column is a single SUM
    over a range — traceable in one click, instead of twelve cells added together.

    The anchor month's column is a formula into Conso_TB_Total rather than a written value, so
    the comparative sheets and PL_Current/BS_Current cannot show different numbers for the
    same month. Every other month is written as data: there is no per-month entity sheet to
    reference, exactly as the Adjustments column has no Entity_ADJ sheet to reference.

    A month with nothing consolidated is left genuinely empty, not zero — the comparative
    sheets then leave that column blank rather than asserting the business did no trade.
    """
    ws = wb.create_sheet("Conso_TB_By_Period")
    months = comparative_engine.fiscal_months()
    current_year, prior_year = context.anchor_year, context.anchor_year - 1

    first_data_col = 4
    col_of = {}
    header = ["Account Code", "Account Name", "Section"]
    for offset, year in enumerate((current_year, prior_year)):
        for i, m in enumerate(months):
            col = first_data_col + offset * len(months) + i
            col_of[comparative_engine.month_key(year, m)] = get_column_letter(col)
            header.append(_month_label(year, m))
    ytd_current_col = get_column_letter(first_data_col + 2 * len(months))
    ytd_prior_col = get_column_letter(first_data_col + 2 * len(months) + 1)
    header += [f"YTD {_month_label(current_year, context.anchor_month)}",
               f"YTD {_month_label(prior_year, context.anchor_month)}"]
    ws.append(header)
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "D2"

    monthly = {}
    for year in (current_year, prior_year):
        for m in months:
            monthly[comparative_engine.month_key(year, m)] = comparative_engine.totals_for_month(conn, year, m)
    has_data = {k for k, v in monthly.items() if v is not None}

    anchor_key = context.anchor_key
    ytd_current_last = col_of[comparative_engine.month_key(current_year, context.anchor_month)]
    ytd_prior_last = col_of[comparative_engine.month_key(prior_year, context.anchor_month)]
    first_col_letter = get_column_letter(first_data_col)
    prior_first_col_letter = get_column_letter(first_data_col + len(months))

    row_of = {}
    for leaf in coa_leaves:
        ws.append([leaf["account_code"], leaf["account_name"], leaf["section"]])
        r = ws.max_row
        row_of[leaf["account_code"]] = r
        for key, totals in monthly.items():
            if totals is None:
                continue
            cell = ws[f"{col_of[key]}{r}"]
            if key == anchor_key:
                cell.value = f"='Conso_TB_Total'!E{code_row[leaf['account_code']]}"
            else:
                entry = totals.get(leaf["account_name"])
                cell.value = entry[0] if entry else 0.0
            cell.number_format = IDR_FORMAT
        ws[f"{ytd_current_col}{r}"] = f"=SUM({first_col_letter}{r}:{ytd_current_last}{r})"
        ws[f"{ytd_prior_col}{r}"] = f"=SUM({prior_first_col_letter}{r}:{ytd_prior_last}{r})"
        for letter in (ytd_current_col, ytd_prior_col):
            ws[f"{letter}{r}"].number_format = IDR_FORMAT

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 20
    for letter in list(col_of.values()) + [ytd_current_col, ytd_prior_col]:
        ws.column_dimensions[letter].width = 15

    return {
        "row": row_of,
        "col": col_of,
        "ytd_current": ytd_current_col,
        "ytd_prior": ytd_prior_col,
        "has_data": has_data,
        "months": months,
        "current_year": current_year,
        "prior_year": prior_year,
    }


def _by_period_ref(name, name_to_leaf, by_period, col_letter, positive_override=False):
    """Leaf reference into one column of Conso_TB_By_Period, with the same Dr-positive ->
    display-magnitude sign flip that _leaf_ref applies for the single-period sheets."""
    info = name_to_leaf[name]
    ref = f"'Conso_TB_By_Period'!{col_letter}{by_period['row'][info['account_code']]}"
    negate = (info["normal_balance"] == "CR") and not positive_override
    return f"-{ref}" if negate else ref


def _comparative_columns(by_period, context, include_ytd):
    """The column plan: YTD pair first (P&L only), then each month of the year immediately
    followed by the same month a year earlier, exactly as the reference workbook orders them.

    Every month gets a column whether or not it has data, so the sheet's shape matches the
    template and stays stable as history is loaded; columns without data are left blank.
    """
    plan = []
    if include_ytd:
        plan.append({"key": comparative_engine.YTD_CURRENT, "source": by_period["ytd_current"],
                     "label": f"YTD {_month_label(by_period['current_year'], context.anchor_month)}",
                     "has_data": any(comparative_engine.month_key(by_period["current_year"], m) in by_period["has_data"]
                                     for m in by_period["months"])})
        plan.append({"key": comparative_engine.YTD_PRIOR, "source": by_period["ytd_prior"],
                     "label": f"YTD {_month_label(by_period['prior_year'], context.anchor_month)}",
                     "has_data": any(comparative_engine.month_key(by_period["prior_year"], m) in by_period["has_data"]
                                     for m in by_period["months"])})
    for m in by_period["months"]:
        for year in (by_period["current_year"], by_period["prior_year"]):
            key = comparative_engine.month_key(year, m)
            plan.append({"key": key, "source": by_period["col"][key], "label": _month_label(year, m),
                         "has_data": key in by_period["has_data"]})
    return plan


def _write_pl_comparative_sheet(wb, conn, name_to_leaf, by_period, context):
    """PL_Comparative: the same P&L rows as PL_Current, one amount-and-% column pair per
    period. Built from the shared PL_LAYOUT, so a subtotal formula cannot differ between this
    sheet and PL_Current."""
    ws = wb.create_sheet("PL_Comparative")
    plan = _comparative_columns(by_period, context, include_ytd=True)

    header, subhead = ["Line Item"], [""]
    for column in plan:
        header += [column["label"], ""]
        subhead += ["IDR", "%"]
    ws.append(header)
    ws.append(subhead)
    for c in ws[1] + ws[2]:
        c.font = BOLD
    ws.freeze_panes = "B3"

    rows = PL_LAYOUT()
    row_of = _place_labels(ws, rows)

    for i, column in enumerate(plan):
        amount_col = get_column_letter(2 + i * 2)
        pct_col = get_column_letter(3 + i * 2)
        ws.column_dimensions[amount_col].width = 17
        ws.column_dimensions[pct_col].width = 9
        ws.cell(row=1, column=2 + i * 2).value = column["label"]
        # A period with nothing consolidated gets no formulas at all. Pointing the leaves at
        # empty cells would render it as a column of zeros, which reads as "we traded nothing"
        # rather than "this month has not been loaded".
        if not column["has_data"]:
            continue
        _fill_column(ws, rows, row_of, amount_col,
                     lambda name, pos, c=column: _by_period_ref(name, name_to_leaf, by_period, c["source"], pos))
        net_sales_row = row_of["Net Sales"]
        for r in row_of.values():
            cell = ws.cell(row=r, column=3 + i * 2)
            cell.value = f"={amount_col}{r}/{amount_col}${net_sales_row}"
            cell.number_format = "0.0%"

    ws.column_dimensions["A"].width = 55
    return row_of["Net Income/(Loss) After Tax & Previous Year Expenses"]


def _write_bs_comparative_sheet(wb, conn, name_to_leaf, by_period, context, pl_comparative_net_income_row):
    """BS_Comparative: paired month-end positions, no year-to-date columns — a balance sheet
    is a point in time and summing one across months is meaningless.

    Each column's Retained Earnings picks up that month's own net income from the matching
    PL_Comparative column, mirroring how BS_Current cross-references PL_Current.
    """
    ws = wb.create_sheet("BS_Comparative")
    plan = _comparative_columns(by_period, context, include_ytd=False)
    pl_plan = _comparative_columns(by_period, context, include_ytd=True)
    pl_col_of = {c["key"]: get_column_letter(2 + i * 2) for i, c in enumerate(pl_plan)}

    ws.append(["Line Item"] + [c["label"] for c in plan])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "B2"

    rows = BS_LAYOUT()
    row_of = _place_labels(ws, rows)

    for i, column in enumerate(plan):
        letter = get_column_letter(2 + i)
        ws.column_dimensions[letter].width = 17
        if not column["has_data"]:
            continue
        ref = lambda name, pos, c=column: _by_period_ref(name, name_to_leaf, by_period, c["source"], pos)
        _fill_column(ws, rows, row_of, letter, ref, ctx={
            "re_leaf": ref("Retained Earnings", False),
            "pl_net_income": lambda col, k=column["key"]: f"'PL_Comparative'!{pl_col_of[k]}{pl_comparative_net_income_row}",
        })

    ws.column_dimensions["A"].width = 55


def _write_adjustments_sheet(wb, conn, period_id):
    from src.db.repositories import adjustment_repo
    ws = wb.create_sheet("Adjustments")
    ws.append(["Adjustment Ref", "Type", "Narration", "Entity", "Category", "Debit", "Credit",
               "Line Narration", "Status", "Copied From"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"
    for adj in adjustment_repo.list_for_period(conn, period_id):
        copied_from = ""
        if adj["copied_from_period_id"]:
            p = conn.execute("SELECT * FROM periods WHERE period_id = ?",
                              (adj["copied_from_period_id"],)).fetchone()
            copied_from = f"{p['year']}-{p['month']:02d}" if p else ""
        for line in adjustment_repo.get_lines(conn, adj["adjustment_id"]):
            entity_code = ""
            if line["entity_id"]:
                e = conn.execute("SELECT entity_code FROM entities WHERE entity_id = ?",
                                  (line["entity_id"],)).fetchone()
                entity_code = e["entity_code"] if e else ""
            category = group_coa_repo.get_by_id(conn, line["group_account_id"])
            ws.append([
                adj["external_ref"], adj["adjustment_type"], adj["narration"], entity_code,
                category["account_name"] if category else "",
                line["debit_amount"] or None, line["credit_amount"] or None,
                line["line_narration"], adj["status"], copied_from,
            ])
            for col in (6, 7):
                ws.cell(row=ws.max_row, column=col).number_format = IDR_FORMAT
    widths = [16, 18, 45, 10, 40, 16, 16, 30, 10, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_adj_bridge_sheet(wb, conn, period_id):
    ws = wb.create_sheet("Adj_Bridge")
    ws.append(["Account Code", "Account Name", "TB Before Adjustments", "Adjustment Impact", "TB After (Total)"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"
    rows = conn.execute(
        """SELECT gc.account_code, gc.account_name,
                  SUM(CASE WHEN ctb.source_type = 'entity_tb' THEN ctb.closing_dr - ctb.closing_cr ELSE 0 END) AS before_adj,
                  SUM(CASE WHEN ctb.source_type = 'adjustment' THEN ctb.closing_dr - ctb.closing_cr ELSE 0 END) AS adj_impact,
                  SUM(CASE WHEN ctb.source_type = 'total' THEN ctb.closing_dr - ctb.closing_cr ELSE 0 END) AS after_total
           FROM consolidated_tb ctb
           JOIN group_coa gc ON gc.group_account_id = ctb.group_account_id
           WHERE ctb.period_id = ?
           GROUP BY ctb.group_account_id
           HAVING adj_impact != 0
           ORDER BY gc.line_item_order""",
        (period_id,),
    ).fetchall()
    for r in rows:
        ws.append([r["account_code"], r["account_name"], r["before_adj"], r["adj_impact"], r["after_total"]])
        for col in (3, 4, 5):
            ws.cell(row=ws.max_row, column=col).number_format = IDR_FORMAT
    for i, w in enumerate([12, 40, 20, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_audit_log_sheet(wb, conn):
    ws = wb.create_sheet("Audit_Log")
    ws.append(["Timestamp", "User", "Action", "Table", "Record ID", "Old Value", "New Value"])
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "A2"
    rows = conn.execute("SELECT * FROM audit_log ORDER BY log_id DESC LIMIT 2000").fetchall()
    for r in rows:
        ws.append([r["timestamp"], r["user"], r["action"], r["table_name"], r["record_id"],
                   r["old_value"], r["new_value"]])
    widths = [24, 16, 12, 22, 12, 30, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_pack(conn: sqlite3.Connection, period_id: int, out_path: str) -> None:
    period = conn.execute("SELECT * FROM periods WHERE period_id = ?", (period_id,)).fetchone()
    period_label = f"{period['year']}-{period['month']:02d}"
    validations = run_all(conn, period_id)

    entities = entity_repo.list_all(conn)
    coa_rows = group_coa_repo.list_all(conn)
    coa_leaves = group_coa_repo.list_leaf_categories(conn)
    name_to_leaf = {r["account_name"]: r for r in coa_leaves}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_summary(wb, conn, period_id, period_label, validations)
    entity_row = _write_entity_sheets(wb, conn, period_id)
    code_row = _write_matrix_and_total(wb, conn, period_id, coa_rows, entities, entity_row)
    net_income_row = _write_pl_sheet(wb, conn, name_to_leaf, code_row)
    _write_bs_sheet(wb, conn, name_to_leaf, code_row, net_income_row)

    context = comparative_engine.resolve_comparative_periods(conn, period_id)
    by_period = _write_by_period_sheet(wb, conn, period_id, coa_leaves, code_row, context)
    comparative_net_income_row = _write_pl_comparative_sheet(wb, conn, name_to_leaf, by_period, context)
    _write_bs_comparative_sheet(wb, conn, name_to_leaf, by_period, context, comparative_net_income_row)

    _write_adjustments_sheet(wb, conn, period_id)
    _write_adj_bridge_sheet(wb, conn, period_id)
    _write_mapping_sheet(wb, conn, period_id)
    _write_validation_sheet(wb, validations)
    _write_audit_log_sheet(wb, conn)
    wb.save(out_path)
