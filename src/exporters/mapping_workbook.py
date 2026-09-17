"""Mapping.xlsx round-trip: export every distinct ledger name for categorization, then
re-import the user's edits. No category is ever pre-filled from anywhere except a prior run of
this same tool (see Phase 1 plan finding #6 — the old Financials workbook's Mapping sheet is
never consulted).

Per-entity splits survive the round-trip. A ledger name booked differently across entities —
different Tally primary group, or already mapped to different categories — exports as a
summary row plus one editable row per entity, and imports through the per-entity write. The
old workbook had one row per ledger name and imported through `apply_mapping`, which writes
one category across every entity carrying that name: re-importing it silently flattened the
split. That is exactly how VKS's shareholder-loan interest was moved out of the P&L into a
balance sheet liability, and nothing in the file told the reader it had happened.

The import is all-or-nothing. Everything is validated first and nothing is written unless all
of it passes, so a rejected file leaves the existing categorization untouched rather than
half-applied.
"""

import sqlite3

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from src.db.repositories import entity_repo, group_coa_repo, mapping_repo, period_repo

HEADER = ["Ledger Name", "Applies To", "Entities Present In", "Tally Primary Group",
          "Balance (IDR)", "Check", "Category"]

# What the "Applies To" cell says for a row that writes one category across every entity.
ALL_ENTITIES = "All entities"
# ...and for a split ledger's summary row, which is not editable.
PER_ENTITY = "(per entity below)"

META_SHEET = "_Meta"


def _period_label(conn, period_id):
    row = conn.execute("SELECT year, month FROM periods WHERE period_id = ?", (period_id,)).fetchone()
    return f"{row['year']}-{row['month']:02d}" if row else ""


def export_mapping_workbook(conn: sqlite3.Connection, period_id: int, out_path: str) -> None:
    ledgers = mapping_repo.ledger_rows_for_period(conn, period_id)
    ledgers.sort(key=lambda x: x["ledger_name"])
    categories = group_coa_repo.list_leaf_categories(conn)
    category_names = [c["account_name"] for c in categories]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"

    editable_rows = []
    for entry in ledgers:
        entities_present = ", ".join(sorted({e["entity_code"] for e in entry["entities"]}))
        groups = sorted({e["primary_group"] for e in entry["entities"] if e["primary_group"]})

        if not entry["needs_split"]:
            ws.append([entry["ledger_name"], ALL_ENTITIES, entities_present, ", ".join(groups),
                       round(entry["total_balance"], 2), "", entry["uniform_category"] or ""])
            editable_rows.append(ws.max_row)
            continue

        # Summary row: names the ledger and says why it is split. Left uneditable so a category
        # typed here cannot be mistaken for one that applied to every entity.
        check = ("Split: differs across entities" if entry["distinct_primary_groups"] > 1
                 else "Split: inconsistently categorized")
        ws.append([entry["ledger_name"], PER_ENTITY, entities_present, ", ".join(groups),
                   round(entry["total_balance"], 2), check, ""])
        for cell in ws[ws.max_row]:
            cell.font = openpyxl.styles.Font(bold=True)

        for e in sorted(entry["entities"], key=lambda x: x["entity_code"]):
            ws.append([entry["ledger_name"], e["entity_code"], e["entity_code"],
                       e["primary_group"], round(e["balance"], 2), "",
                       e["category_name"] or ""])
            editable_rows.append(ws.max_row)

    cat_sheet = wb.create_sheet("_CategoryList")
    for i, name in enumerate(category_names, start=1):
        cat_sheet.cell(row=i, column=1, value=name)
    cat_sheet.sheet_state = "hidden"

    # The period the file was exported for, so the import can work out which ledgers are split
    # without the caller having to remember and pass it back in.
    meta = wb.create_sheet(META_SHEET)
    meta["A1"] = "period"
    meta["B1"] = _period_label(conn, period_id)
    meta.sheet_state = "hidden"

    dv = DataValidation(
        type="list",
        formula1=f"'_CategoryList'!$A$1:$A${len(category_names)}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid category",
        error="Pick a category from the dropdown list.",
    )
    ws.add_data_validation(dv)
    for r in editable_rows:
        dv.add(f"G{r}")

    widths = [40, 18, 20, 28, 16, 30, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(out_path)


def _resolve_period(conn, wb, period_id):
    """The period the file belongs to: the caller's, else the one stamped at export, else the
    most recent period holding a TB import."""
    if period_id is not None:
        return period_id
    if META_SHEET in wb.sheetnames:
        label = wb[META_SHEET]["B1"].value
        if label:
            period = period_repo.get_by_str(conn, str(label).strip())
            if period:
                return period["period_id"]
    row = conn.execute(
        """SELECT p.period_id FROM periods p
           JOIN trial_balance_imports tbi ON tbi.period_id = p.period_id
           GROUP BY p.period_id ORDER BY p.year DESC, p.month DESC LIMIT 1"""
    ).fetchone()
    return row["period_id"] if row else None


def import_mapping_workbook(conn: sqlite3.Connection, in_path: str, user: str | None = None,
                            period_id: int | None = None) -> int:
    wb = openpyxl.load_workbook(in_path, data_only=True)
    ws = wb["Mapping"]

    categories = group_coa_repo.list_leaf_categories(conn)
    code_by_name = {c["account_name"]: c["group_account_id"] for c in categories}
    entity_by_code = {e["entity_code"]: e["entity_id"] for e in entity_repo.list_all(conn)}

    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header) if name}
    required = {"Ledger Name", "Category"}
    if not required.issubset(col):
        raise ValueError(f"Mapping.xlsx is missing required column(s): {required - set(col)}")

    resolved_period = _resolve_period(conn, wb, period_id)
    # Recomputed from the database, never read from the file: a ledger can have become split
    # since the workbook was exported, and the file would not know.
    split_ledgers = set()
    if resolved_period is not None:
        split_ledgers = {r["ledger_name"] for r in mapping_repo.ledger_rows_for_period(conn, resolved_period)
                         if r["needs_split"]}

    pending = []
    errors = []
    for row in range(2, ws.max_row + 1):
        ledger_name = ws.cell(row=row, column=col["Ledger Name"]).value
        if not ledger_name:
            continue
        ledger_name = str(ledger_name).strip()
        category = ws.cell(row=row, column=col["Category"]).value
        applies_to = ws.cell(row=row, column=col["Applies To"]).value if "Applies To" in col else None
        applies_to = str(applies_to).strip() if applies_to else ALL_ENTITIES

        if applies_to == PER_ENTITY:
            # The summary row above a split ledger's per-entity rows. Blank is the normal case;
            # a category typed here is refused rather than ignored, so nobody believes it
            # applied to every entity when it did not.
            if category:
                errors.append(
                    f"Row {row}: '{ledger_name}' is split across entities — set the category on "
                    f"its per-entity rows, not on the {PER_ENTITY} summary row")
            continue

        if not category:
            continue  # still uncategorized, nothing to apply
        category = str(category).strip()
        if category not in code_by_name:
            errors.append(f"Row {row}: unknown category '{category}' for ledger '{ledger_name}'")
            continue

        if applies_to == ALL_ENTITIES:
            if ledger_name in split_ledgers:
                errors.append(
                    f"Row {row}: '{ledger_name}' is booked differently across entities, so one "
                    f"category cannot be applied to all of them. Re-export the mapping workbook "
                    f"to get a row per entity.")
                continue
            pending.append((ledger_name, None, code_by_name[category]))
        elif applies_to in entity_by_code:
            pending.append((ledger_name, entity_by_code[applies_to], code_by_name[category]))
        else:
            errors.append(f"Row {row}: unknown entity '{applies_to}' for ledger '{ledger_name}'")

    if errors:
        raise ValueError("Mapping.xlsx was not applied:\n" + "\n".join(errors))

    # Nothing is written until every row has passed, so a rejected file leaves the existing
    # categorization exactly as it was rather than half-updated.
    for ledger_name, entity_id, group_account_id in pending:
        if entity_id is None:
            mapping_repo.apply_mapping(conn, ledger_name, group_account_id, user=user)
        else:
            mapping_repo.apply_mapping_for_entity(conn, entity_id, ledger_name, group_account_id, user=user)

    return len(pending)
