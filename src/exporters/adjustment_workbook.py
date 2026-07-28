"""Adjustments.xlsx round-trip — JE-style grid, one row per adjustment line, grouped by
a user-assigned Adjustment Ref (matches the plan's "Adjustment entry form (JE-style grid)"
UI screen, translated into this project's CLI + Excel round-trip pattern)."""

import sqlite3

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from src.db.repositories import adjustment_repo, entity_repo, group_coa_repo
from src.db.repositories.adjustment_repo import VALID_TYPES
from src.engines.adjustment_engine import create_adjustment

HEADER = ["Adjustment Ref", "Type", "Narration", "Entity Code", "Category", "Debit", "Credit",
          "Line Narration", "Status"]


def export_adjustment_workbook(conn: sqlite3.Connection, period_id: int, out_path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Adjustments"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)
    ws.freeze_panes = "A2"

    for adj in adjustment_repo.list_for_period(conn, period_id):
        for line in adjustment_repo.get_lines(conn, adj["adjustment_id"]):
            entity_code = ""
            if line["entity_id"]:
                e = conn.execute("SELECT entity_code FROM entities WHERE entity_id = ?",
                                  (line["entity_id"],)).fetchone()
                entity_code = e["entity_code"] if e else ""
            category = group_coa_repo.get_by_id(conn, line["group_account_id"])
            ws.append([
                adj["external_ref"], adj["adjustment_type"], adj["narration"], entity_code,
                category["account_name"] if category else "",
                line["debit_amount"] or None, line["credit_amount"] or None,
                line["line_narration"], adj["status"],
            ])

    categories = group_coa_repo.list_leaf_categories(conn)
    entities = entity_repo.list_all(conn)
    cat_sheet = wb.create_sheet("_CategoryList")
    for i, c in enumerate(categories, start=1):
        cat_sheet.cell(row=i, column=1, value=c["account_name"])
    cat_sheet.sheet_state = "hidden"
    ent_sheet = wb.create_sheet("_EntityList")
    ent_sheet.cell(row=1, column=1, value="")  # blank = group-level line
    for i, e in enumerate(entities, start=2):
        ent_sheet.cell(row=i, column=1, value=e["entity_code"])
    ent_sheet.sheet_state = "hidden"
    type_sheet = wb.create_sheet("_TypeList")
    for i, t in enumerate(sorted(VALID_TYPES), start=1):
        type_sheet.cell(row=i, column=1, value=t)
    type_sheet.sheet_state = "hidden"

    max_row = max(ws.max_row, 200)
    dv_cat = DataValidation(type="list", formula1=f"'_CategoryList'!$A$1:$A${len(categories)}", allow_blank=True)
    ws.add_data_validation(dv_cat)
    dv_cat.add(f"E2:E{max_row}")
    dv_ent = DataValidation(type="list", formula1=f"'_EntityList'!$A$1:$A${len(entities) + 1}", allow_blank=True)
    ws.add_data_validation(dv_ent)
    dv_ent.add(f"D2:D{max_row}")
    dv_type = DataValidation(type="list", formula1=f"'_TypeList'!$A$1:$A${len(VALID_TYPES)}", allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(f"B2:B{max_row}")

    widths = [16, 18, 45, 12, 40, 16, 16, 30, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(out_path)


def import_adjustment_workbook(conn: sqlite3.Connection, period_id: int, in_path: str,
                                 user: str | None = None) -> int:
    wb = openpyxl.load_workbook(in_path, data_only=True)
    ws = wb["Adjustments"]
    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}
    required = {"Adjustment Ref", "Type", "Narration", "Category", "Debit", "Credit"}
    if not required.issubset(col):
        raise ValueError(f"Adjustments.xlsx is missing required column(s): {required - set(col)}")

    grouped: dict[str, dict] = {}
    for row in range(2, ws.max_row + 1):
        ref = ws.cell(row=row, column=col["Adjustment Ref"]).value
        if not ref:
            continue
        ref = str(ref).strip()
        category_name = ws.cell(row=row, column=col["Category"]).value
        if not category_name:
            continue
        category = group_coa_repo.get_by_name(conn, str(category_name).strip())
        if category is None:
            raise ValueError(f"Row {row}: unknown category '{category_name}'")

        entity_code = ws.cell(row=row, column=col["Entity Code"]).value
        entity_id = None
        if entity_code:
            e = entity_repo.get_by_code(conn, str(entity_code).strip())
            if e is None:
                raise ValueError(f"Row {row}: unknown entity code '{entity_code}'")
            entity_id = e["entity_id"]

        entry = grouped.setdefault(ref, {
            "adjustment_type": (ws.cell(row=row, column=col["Type"]).value or "other").strip(),
            "narration": (ws.cell(row=row, column=col["Narration"]).value or "").strip(),
            "lines": [],
        })
        entry["lines"].append({
            "entity_id": entity_id,
            "group_account_id": category["group_account_id"],
            "debit_amount": float(ws.cell(row=row, column=col["Debit"]).value or 0),
            "credit_amount": float(ws.cell(row=row, column=col["Credit"]).value or 0),
            "line_narration": ws.cell(row=row, column=col.get("Line Narration", 0)).value
                if "Line Narration" in col else None,
        })

    for ref, entry in grouped.items():
        create_adjustment(conn, period_id, ref, entry["adjustment_type"], entry["narration"],
                            entry["lines"], user=user)

    return len(grouped)
