import argparse
import getpass
import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
CONFIG_DIR = PROJECT_ROOT / "config"

from src.db.database import get_connection
from src.db.seed import seed_all
from src.db.repositories import entity_repo, period_repo, tb_repo, mapping_repo
from src.importers.tally_tb_parser import parse_file


def _settings() -> dict:
    return json.loads((CONFIG_DIR / "app_settings.json").read_text(encoding="utf-8"))


def _db_path() -> str:
    """database_path in config is relative — resolve it against PROJECT_ROOT, not the
    process's CWD, so `python -m src.cli ...` behaves the same regardless of where it's
    invoked from (see src/web/deps.py's _db_path for the concrete bug this guards against)."""
    raw = _settings()["database_path"]
    path = Path(raw)
    return str(path if path.is_absolute() else PROJECT_ROOT / path)


def _connect():
    conn = get_connection(_db_path())
    seed_all(conn)  # idempotent: re-applies fixed entities/group_coa, never touches mappings/TBs
    return conn


def cmd_import_tb(args: argparse.Namespace) -> None:
    conn = _connect()
    entity = entity_repo.get_by_code(conn, args.entity)
    if not entity:
        sys.exit(f"Unknown entity code '{args.entity}'. Check config/app_settings.json.")
    period_id = period_repo.get_or_create(conn, *period_repo.parse_period_str(args.period))

    parsed = parse_file(args.file)
    if tb_repo.find_duplicate_checksum(conn, period_id, entity["entity_id"], parsed.checksum):
        print(f"Note: identical file already imported for {args.entity} {args.period} "
              f"(checksum match) — re-importing anyway.")

    try:
        import_id = tb_repo.import_tb(conn, period_id, entity["entity_id"], parsed, imported_by=getpass.getuser())
    except ValueError as e:
        sys.exit(str(e))
    mapping_repo.ensure_rows_exist(conn, period_id)
    conn.commit()

    print(f"Imported {args.file}")
    print(f"  entity={args.entity}  period={args.period}  import_id={import_id}")
    print(f"  lines={len(parsed.lines)}  grand_total_dr={parsed.grand_total_period_dr:,.2f}  "
          f"grand_total_cr={parsed.grand_total_period_cr:,.2f}")
    if parsed.warnings:
        print(f"  {len(parsed.warnings)} warning(s):")
        for w in parsed.warnings:
            print(f"    ! {w}")


def cmd_export_mapping(args: argparse.Namespace) -> None:
    from src.exporters.mapping_workbook import export_mapping_workbook
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    mapping_repo.ensure_rows_exist(conn, period["period_id"])
    conn.commit()
    export_mapping_workbook(conn, period["period_id"], args.out)
    print(f"Wrote {args.out}")


def cmd_import_mapping(args: argparse.Namespace) -> None:
    from src.exporters.mapping_workbook import import_mapping_workbook
    conn = _connect()
    n = import_mapping_workbook(conn, args.in_file, user=getpass.getuser())
    conn.commit()
    print(f"Applied categorization to {n} ledger(s).")


def cmd_consolidate(args: argparse.Namespace) -> None:
    from src.engines.consolidation_engine import consolidate_period
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    result = consolidate_period(conn, period["period_id"])
    conn.commit()
    if result["blocked"]:
        if result.get("reason") == "locked":
            sys.exit(f"BLOCKED: period {args.period} is locked. Unlock it first if you really need to re-consolidate.")
        print(f"BLOCKED: {len(result['unmapped_material'])} ledger(s) with a non-zero balance "
              f"are still uncategorized. Consolidation was not run.")
        for name, bal in result["unmapped_material"][:30]:
            print(f"    - {name} ({bal:,.2f})")
        if len(result["unmapped_material"]) > 30:
            print(f"    ... and {len(result['unmapped_material']) - 30} more")
        sys.exit(1)
    print(f"Consolidated period {args.period}: {result['rows_written']} consolidated_tb rows.")


def cmd_export_pack(args: argparse.Namespace) -> None:
    from src.exporters.excel_pack_generator import export_pack
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    export_pack(conn, period["period_id"], args.out)
    print(f"Wrote {args.out}")


def cmd_export_adjustments(args: argparse.Namespace) -> None:
    from src.exporters.adjustment_workbook import export_adjustment_workbook
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    export_adjustment_workbook(conn, period["period_id"], args.out)
    print(f"Wrote {args.out}")


def cmd_import_adjustments(args: argparse.Namespace) -> None:
    from src.exporters.adjustment_workbook import import_adjustment_workbook
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    try:
        n = import_adjustment_workbook(conn, period["period_id"], args.in_file, user=getpass.getuser())
    except ValueError as e:
        sys.exit(str(e))
    conn.commit()
    print(f"Imported {n} adjustment(s) as draft.")


def cmd_apply_adjustments(args: argparse.Namespace) -> None:
    from src.db.repositories import adjustment_repo
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    try:
        n = adjustment_repo.apply_all(conn, period["period_id"], user=getpass.getuser())
    except ValueError as e:
        sys.exit(str(e))
    conn.commit()
    print(f"Applied {n} draft adjustment(s). Re-run consolidate to flow them into the consolidated TB.")


def cmd_copy_adjustments(args: argparse.Namespace) -> None:
    from src.engines.adjustment_engine import copy_prior_month
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    try:
        n = copy_prior_month(conn, period["period_id"], user=getpass.getuser())
    except ValueError as e:
        sys.exit(str(e))
    conn.commit()
    print(f"Copied {n} adjustment(s) from the prior period as new drafts.")


def cmd_delete_adjustment(args: argparse.Namespace) -> None:
    from src.db.repositories import adjustment_repo
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found.")
    row = conn.execute(
        "SELECT adjustment_id FROM adjustments WHERE period_id = ? AND external_ref = ?",
        (period["period_id"], args.ref),
    ).fetchone()
    if not row:
        sys.exit(f"No adjustment '{args.ref}' found in period {args.period}.")
    try:
        adjustment_repo.delete(conn, period["period_id"], row["adjustment_id"], user=getpass.getuser())
    except ValueError as e:
        sys.exit(str(e))
    conn.commit()
    print(f"Deleted adjustment {args.ref}. Re-run consolidate to remove its effect from the consolidated TB.")


def cmd_edit_adjustment(args: argparse.Namespace) -> None:
    """Loads only --ref's rows from the workbook and routes them through update_adjustment
    — the single-entry, audited path. import-adjustments already round-trips everything;
    this is for editing one existing entry with a before/after audit trail."""
    import openpyxl
    from src.db.repositories import group_coa_repo, adjustment_repo
    from src.engines.adjustment_engine import update_adjustment

    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")

    adj = conn.execute(
        "SELECT * FROM adjustments WHERE period_id = ? AND external_ref = ?",
        (period["period_id"], args.ref),
    ).fetchone()
    if adj is None:
        sys.exit(f"No adjustment '{args.ref}' found in period {args.period}.")

    wb = openpyxl.load_workbook(args.in_file, data_only=True)
    ws = wb["Adjustments"]
    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}
    required = {"Adjustment Ref", "Type", "Narration", "Category", "Debit", "Credit"}
    if not required.issubset(col):
        sys.exit(f"{args.in_file} is missing required column(s): {required - set(col)}")

    adjustment_type = adj["adjustment_type"]
    narration = adj["narration"]
    lines = []
    for row in range(2, ws.max_row + 1):
        ref = ws.cell(row=row, column=col["Adjustment Ref"]).value
        if not ref or str(ref).strip() != args.ref:
            continue
        category_name = ws.cell(row=row, column=col["Category"]).value
        if not category_name:
            continue
        category = group_coa_repo.get_by_name(conn, str(category_name).strip())
        if category is None:
            sys.exit(f"Row {row}: unknown category '{category_name}'")

        entity_code = ws.cell(row=row, column=col.get("Entity Code", 0)).value if "Entity Code" in col else None
        entity_id = None
        if entity_code:
            e = entity_repo.get_by_code(conn, str(entity_code).strip())
            if e is None:
                sys.exit(f"Row {row}: unknown entity code '{entity_code}'")
            entity_id = e["entity_id"]

        adjustment_type = (ws.cell(row=row, column=col["Type"]).value or adjustment_type).strip()
        narration = (ws.cell(row=row, column=col["Narration"]).value or narration).strip()
        lines.append({
            "entity_id": entity_id,
            "group_account_id": category["group_account_id"],
            "debit_amount": float(ws.cell(row=row, column=col["Debit"]).value or 0),
            "credit_amount": float(ws.cell(row=row, column=col["Credit"]).value or 0),
            "line_narration": ws.cell(row=row, column=col.get("Line Narration", 0)).value
                if "Line Narration" in col else None,
        })

    if not lines:
        sys.exit(f"No rows for ref '{args.ref}' found in {args.in_file}.")

    try:
        update_adjustment(conn, period["period_id"], adj["adjustment_id"], adjustment_type, narration,
                            lines, user=getpass.getuser())
    except ValueError as e:
        sys.exit(str(e))
    conn.commit()
    print(f"Updated {args.ref} — back to draft. Apply it and re-run consolidate.")


def cmd_lock_period(args: argparse.Namespace) -> None:
    from src.engines.validation_engine import run_all
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found — run import-tb first.")
    results = run_all(conn, period["period_id"])
    errors = [r for r in results if not r.passed and r.severity == "Error"]
    if errors:
        print("BLOCKED: cannot lock — failing validations:")
        for r in errors:
            print(f"    - {r.check_id}: {r.description}")
        sys.exit(1)
    period_repo.lock(conn, period["period_id"], user=getpass.getuser())
    conn.commit()
    print(f"Locked period {args.period}.")


def cmd_unlock_period(args: argparse.Namespace) -> None:
    conn = _connect()
    period = period_repo.get_by_str(conn, args.period)
    if not period:
        sys.exit(f"No period '{args.period}' found.")
    period_repo.unlock(conn, period["period_id"], user=getpass.getuser())
    conn.commit()
    print(f"Unlocked period {args.period}.")


def main() -> None:
    parser = argparse.ArgumentParser(prog="tally-consolidation")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("import-tb", help="Import a Tally TB Excel export for one entity/period")
    p.add_argument("--entity", required=True, help="Entity code, e.g. BII")
    p.add_argument("--period", required=True, help="Period as YYYY-MM, e.g. 2026-06")
    p.add_argument("file", help="Path to the Tally TB Excel export")
    p.set_defaults(func=cmd_import_tb)

    p = sub.add_parser("export-mapping", help="Export the ledger categorization workbook")
    p.add_argument("--period", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=cmd_export_mapping)

    p = sub.add_parser("import-mapping", help="Re-import an edited categorization workbook")
    p.add_argument("--in", dest="in_file", required=True)
    p.set_defaults(func=cmd_import_mapping)

    p = sub.add_parser("consolidate", help="Roll up mapped entity TBs into the consolidated TB")
    p.add_argument("--period", required=True)
    p.set_defaults(func=cmd_consolidate)

    p = sub.add_parser("export-pack", help="Export the consolidated Excel pack")
    p.add_argument("--period", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=cmd_export_pack)

    p = sub.add_parser("export-adjustments", help="Export the adjustments workbook")
    p.add_argument("--period", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(func=cmd_export_adjustments)

    p = sub.add_parser("import-adjustments", help="Re-import an edited adjustments workbook (as drafts)")
    p.add_argument("--period", required=True)
    p.add_argument("--in", dest="in_file", required=True)
    p.set_defaults(func=cmd_import_adjustments)

    p = sub.add_parser("edit-adjustment", help="Edit a single existing adjustment from a workbook, with an audited before/after trail")
    p.add_argument("--period", required=True)
    p.add_argument("--ref", required=True, help="external_ref of the adjustment to edit")
    p.add_argument("--in", dest="in_file", required=True)
    p.set_defaults(func=cmd_edit_adjustment)

    p = sub.add_parser("apply-adjustments", help="Apply all draft adjustments for a period")
    p.add_argument("--period", required=True)
    p.set_defaults(func=cmd_apply_adjustments)

    p = sub.add_parser("copy-adjustments", help="Copy the prior period's adjustments in as new drafts")
    p.add_argument("--period", required=True, help="Target period to copy into")
    p.set_defaults(func=cmd_copy_adjustments)

    p = sub.add_parser("delete-adjustment", help="Delete a single adjustment by its ref (e.g. to remove a test entry)")
    p.add_argument("--period", required=True)
    p.add_argument("--ref", required=True, help="The Adjustment Ref to delete")
    p.set_defaults(func=cmd_delete_adjustment)

    p = sub.add_parser("lock-period", help="Lock a period (requires all Error-severity validations to pass)")
    p.add_argument("--period", required=True)
    p.set_defaults(func=cmd_lock_period)

    p = sub.add_parser("unlock-period", help="Unlock a period (admin override, logged)")
    p.add_argument("--period", required=True)
    p.set_defaults(func=cmd_unlock_period)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
