"""Regression tests for the Excel pack's formula chain.

The bug these exist to catch: a GROUP-LEVEL adjustment (adjustment_lines.entity_id IS NULL)
belongs to no entity column. It was being left out of Conso_TB_Matrix's Total column, which
summed only the entity columns, so the exported Balance Sheet silently disagreed with the
tool. Nothing caught it because a balanced group-level journal leaves Total Assets, Total
Liabilities and Equity, and the Check row all still tying.

There is no Excel here to recalculate the workbook, so these tests evaluate the formula
chain themselves (Entity_<CODE> -> Conso_TB_Matrix -> Conso_TB_Total -> PL_Current/BS_Current)
and compare the result against the engine that drives the on-screen statements.
"""

import re

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string, get_column_letter

from src.db.repositories import adjustment_repo, consolidated_repo, group_coa_repo
from src.engines.adjustment_engine import create_adjustment
from src.engines.consolidation_engine import consolidate_period
from src.engines.statement_generator import compute_bs, compute_pl
from src.exporters.excel_pack_generator import export_pack

TOLERANCE = 1.0


# --------------------------------------------------------------------------- evaluator

_SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z_]\w*))!\$?([A-Z]{1,3})\$?(\d+)")
_LOCAL_REF = re.compile(r"(?<![A-Z0-9!$:])\$?([A-Z]{1,3})\$?(\d+)(?![(\d])")
_SUM = re.compile(r"SUM\(([^)]*)\)", re.I)
_ENDPOINT = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_]\w*))!)?\$?([A-Z]{1,3})\$?(\d+)")


class Workbook:
    """Just enough of Excel to resolve this pack: cross-sheet refs, SUM(range), + - ( )."""

    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path, data_only=False)
        self._cache = {}

    def cell(self, sheet, col, row):
        key = (sheet, col, row)
        if key in self._cache:
            return self._cache[key]
        self._cache[key] = 0.0  # guards against a circular reference hanging the test
        raw = self.wb[sheet].cell(row=row, column=column_index_from_string(col)).value
        self._cache[key] = self._eval(raw, sheet)
        return self._cache[key]

    def _sum_range(self, expr, ctx):
        start, end = expr.split(':')
        m = _ENDPOINT.match(start)
        sheet = m.group(1) or m.group(2) or ctx
        c1, r1 = m.group(3), int(m.group(4))
        m = _ENDPOINT.match(end)
        c2, r2 = m.group(3), int(m.group(4))
        return sum(
            self.cell(sheet, get_column_letter(ci), ri)
            for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1)
            for ri in range(r1, r2 + 1)
        )

    def _eval(self, raw, ctx):
        if isinstance(raw, (int, float)):
            return float(raw)
        if not isinstance(raw, str) or not raw.startswith('='):
            return 0.0
        expr = raw[1:]
        expr = _SUM.sub(lambda m: repr(self._sum_range(m.group(1).strip(), ctx)), expr)
        expr = _SHEET_REF.sub(
            lambda m: repr(self.cell(m.group(1) or m.group(2), m.group(3), int(m.group(4)))), expr)
        expr = _LOCAL_REF.sub(lambda m: repr(self.cell(ctx, m.group(1), int(m.group(2)))), expr)
        return float(eval(expr, {"__builtins__": {}}, {}))

    def labelled(self, sheet, label_col=1, value_col=2):
        """[(label, evaluated value)] for every row carrying a value."""
        ws = self.wb[sheet]
        out = []
        for r in range(2, ws.max_row + 1):
            label = ws.cell(r, label_col).value
            raw = ws.cell(r, value_col).value
            if label is None or raw is None:
                continue
            out.append((str(label).strip(), self._eval(raw, sheet)))
        return out

    def row_of(self, sheet, account_code):
        ws = self.wb[sheet]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == account_code:
                return r
        raise AssertionError(f"{account_code} not found in {sheet}")


# --------------------------------------------------------------------------- helpers

def _account(conn, name):
    row = group_coa_repo.get_by_name(conn, name)
    assert row is not None, f"unknown category {name}"
    return row["group_account_id"]


def _book(conn, period_id, ref, lines, adj_type="reclassification"):
    """Create + apply an adjustment, then re-consolidate."""
    create_adjustment(conn, period_id, ref, adj_type, f"{ref} test entry", lines, user="test")
    adjustment_repo.apply_all(conn, period_id, user="test")
    conn.commit()
    result = consolidate_period(conn, period_id)
    conn.commit()
    assert not result["blocked"], result
    return result


def _export(conn, period_id, tmp_path, name="pack.xlsx"):
    path = tmp_path / name
    export_pack(conn, period_id, str(path))
    return Workbook(str(path))


GROUP_AMOUNT = 1_234_567_890.0
ENTITY_AMOUNT = 987_654_321.0


@pytest.fixture
def group_level_booked(consolidated, tmp_path):
    """A balanced adjustment with NO entity on either line — the case that was being lost."""
    conn, period_id, _ = consolidated
    _book(conn, period_id, "ADJ-GROUP", [
        {"entity_id": None, "group_account_id": _account(conn, "Accounts Payable"),
         "debit_amount": GROUP_AMOUNT, "credit_amount": 0.0},
        {"entity_id": None, "group_account_id": _account(conn, "Other Payables"),
         "debit_amount": 0.0, "credit_amount": GROUP_AMOUNT},
    ])
    return conn, period_id, _export(conn, period_id, tmp_path)


# --------------------------------------------------------------------------- the bug

def test_group_level_adjustment_reaches_the_matrix_total(group_level_booked):
    conn, period_id, wb = group_level_booked
    _, group_adjustment, total_by_account = consolidated_repo.get_buckets(conn, period_id)

    ap_id = _account(conn, "Accounts Payable")
    assert group_adjustment[ap_id] == pytest.approx(GROUP_AMOUNT), \
        "the adjustment should be bucketed as group-level, not against an entity"

    row = wb.row_of("Conso_TB_Matrix", "BS015")
    matrix_total = wb.cell("Conso_TB_Matrix", get_column_letter(wb.wb["Conso_TB_Matrix"].max_column), row)
    assert matrix_total == pytest.approx(total_by_account[ap_id], abs=TOLERANCE), \
        "Conso_TB_Matrix's Total must include group-level adjustments"


def test_group_level_adjustment_reaches_the_excel_balance_sheet(group_level_booked):
    conn, period_id, wb = group_level_booked
    tool = {l.label: l.value for l in compute_bs(conn, period_id)[0]}
    excel = dict(wb.labelled("BS_Current"))

    for line in ("Accounts Payable", "Other Payables", "Total Short Term Liabilities"):
        assert excel[line] == pytest.approx(tool[line], abs=TOLERANCE), f"{line} disagrees"


def test_every_balance_sheet_line_matches_the_tool(group_level_booked):
    conn, period_id, wb = group_level_booked
    tool = {l.label: l.value for l in compute_bs(conn, period_id)[0] if l.value is not None}
    mismatched = [
        (label, value, tool[label])
        for label, value in wb.labelled("BS_Current")
        if label in tool and abs(value - tool[label]) > TOLERANCE
    ]
    assert not mismatched, f"Excel disagrees with compute_bs on: {mismatched}"


def test_every_pl_line_matches_the_tool(group_level_booked):
    conn, period_id, wb = group_level_booked
    tool = {l.label: l.value for l in compute_pl(conn, period_id)[0] if l.value is not None}
    mismatched = [
        (label, value, tool[label])
        for label, value in wb.labelled("PL_Current")
        if label in tool and abs(value - tool[label]) > TOLERANCE
    ]
    assert not mismatched, f"Excel disagrees with compute_pl on: {mismatched}"


def test_a_balanced_group_journal_does_not_hide_in_the_totals(group_level_booked):
    """Why this went unnoticed: the dropped entry was a balanced journal inside
    liabilities+equity, so these three tied even while six lines were wrong. They must
    still tie AFTER the fix — otherwise the fix broke the workbook a different way."""
    _, _, wb = group_level_booked
    excel = dict(wb.labelled("BS_Current"))
    assert excel["TOTAL LIABILITIES AND EQUITY"] == pytest.approx(excel["TOTAL ASSETS"], abs=TOLERANCE)
    assert excel["Check"] == pytest.approx(0.0, abs=TOLERANCE)


# ------------------------------------------------- the mirror-image bug on the same cause

def test_consolidated_tb_total_does_not_double_count_group_adjustments(group_level_booked):
    """Group-level adjustment rows and 'total' rows BOTH carry entity_id IS NULL. Bucketing
    by entity_id alone merged them and double-counted every group-level adjustment."""
    conn, period_id, _ = group_level_booked
    _, _, total_by_account = consolidated_repo.get_buckets(conn, period_id)
    ap_id = _account(conn, "Accounts Payable")

    authoritative = conn.execute(
        """SELECT SUM(closing_dr - closing_cr) FROM consolidated_tb
           WHERE period_id = ? AND group_account_id = ? AND source_type = 'total'""",
        (period_id, ap_id),
    ).fetchone()[0]
    assert total_by_account[ap_id] == pytest.approx(authoritative, abs=TOLERANCE)


def test_entity_columns_exclude_group_adjustments(group_level_booked):
    conn, period_id, wb = group_level_booked
    by_account_entity, _, _ = consolidated_repo.get_buckets(conn, period_id)
    ap_id = _account(conn, "Accounts Payable")
    assert (ap_id, None) not in by_account_entity, \
        "a group-level adjustment must never land in an entity bucket"


# -------------------------------------------------- entity-specific must NOT be doubled

def test_entity_specific_adjustment_is_not_double_counted(consolidated, tmp_path):
    """The other half of the fix: entity-specific adjustments are already inside the
    Entity_<CODE> sheets, so the Adjustments column must not repeat them."""
    conn, period_id, _ = consolidated
    kns = conn.execute("SELECT entity_id FROM entities WHERE entity_code = 'KNS'").fetchone()["entity_id"]
    _book(conn, period_id, "ADJ-ENTITY", [
        {"entity_id": kns, "group_account_id": _account(conn, "Accounts Payable"),
         "debit_amount": ENTITY_AMOUNT, "credit_amount": 0.0},
        {"entity_id": kns, "group_account_id": _account(conn, "Other Payables"),
         "debit_amount": 0.0, "credit_amount": ENTITY_AMOUNT},
    ])
    wb = _export(conn, period_id, tmp_path)

    ws = wb.wb["Conso_TB_Matrix"]
    row = wb.row_of("Conso_TB_Matrix", "BS015")
    adj_cell = ws.cell(row=row, column=ws.max_column - 1).value
    assert adj_cell in (None, 0), "entity-specific adjustments must not appear in the Adjustments column"

    tool = {l.label: l.value for l in compute_bs(conn, period_id)[0]}
    excel = dict(wb.labelled("BS_Current"))
    assert excel["Accounts Payable"] == pytest.approx(tool["Accounts Payable"], abs=TOLERANCE)


def test_mixed_entity_and_group_adjustments_on_one_account(consolidated, tmp_path):
    """Both kinds hitting the same account — the arrangement that produced the real
    August discrepancy on Retained Earnings."""
    conn, period_id, _ = consolidated
    bii = conn.execute("SELECT entity_id FROM entities WHERE entity_code = 'BII'").fetchone()["entity_id"]
    _book(conn, period_id, "ADJ-MIXED-ENTITY", [
        {"entity_id": bii, "group_account_id": _account(conn, "Retained Earnings"),
         "debit_amount": ENTITY_AMOUNT, "credit_amount": 0.0},
        {"entity_id": bii, "group_account_id": _account(conn, "Inventory"),
         "debit_amount": 0.0, "credit_amount": ENTITY_AMOUNT},
    ])
    _book(conn, period_id, "ADJ-MIXED-GROUP", [
        {"entity_id": None, "group_account_id": _account(conn, "Retained Earnings"),
         "debit_amount": GROUP_AMOUNT, "credit_amount": 0.0},
        {"entity_id": None, "group_account_id": _account(conn, "Loans & Advances taken"),
         "debit_amount": 0.0, "credit_amount": GROUP_AMOUNT},
    ])
    wb = _export(conn, period_id, tmp_path)

    tool = {l.label: l.value for l in compute_bs(conn, period_id)[0] if l.value is not None}
    mismatched = [
        (label, value, tool[label])
        for label, value in wb.labelled("BS_Current")
        if label in tool and abs(value - tool[label]) > TOLERANCE
    ]
    assert not mismatched, f"Excel disagrees with compute_bs on: {mismatched}"


def test_pack_with_no_adjustments_still_ties(consolidated, tmp_path):
    conn, period_id, _ = consolidated
    wb = _export(conn, period_id, tmp_path)
    tool = {l.label: l.value for l in compute_bs(conn, period_id)[0] if l.value is not None}
    mismatched = [
        (label, value, tool[label])
        for label, value in wb.labelled("BS_Current")
        if label in tool and abs(value - tool[label]) > TOLERANCE
    ]
    assert not mismatched, f"Excel disagrees with compute_bs on: {mismatched}"
