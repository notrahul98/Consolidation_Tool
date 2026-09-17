"""Mapping.xlsx must not flatten a per-entity split (Phase 8.0).

The old workbook had one row per ledger name and imported through `apply_mapping`, which
writes a single category across every entity carrying that name. Exporting and re-importing an
untouched file was therefore destructive for any ledger booked differently across entities —
and destructive silently, because the file gave the reader no hint that a split existed. That
is how VKS's shareholder-loan interest ended up as a balance sheet liability instead of a P&L
expense.
"""

import openpyxl
import pytest

from src.db.repositories import group_coa_repo, mapping_repo
from src.exporters.mapping_workbook import (
    ALL_ENTITIES,
    HEADER,
    PER_ENTITY,
    export_mapping_workbook,
    import_mapping_workbook,
)
from tests.conftest import import_all_entities

LEDGER = "12% Interest on Shareholders Loan A/c"


def _entity_id(conn, code):
    return conn.execute("SELECT entity_id FROM entities WHERE entity_code = ?", (code,)).fetchone()["entity_id"]


def _category_for(conn, code, ledger=LEDGER):
    row = mapping_repo.get_mapping(conn, _entity_id(conn, code), ledger)
    return row["group_account_id"] if row else None


def _split_the_interest_ledger(conn):
    """Put the shareholder-loan interest ledger into the state that caused the original bug:
    a P&L expense for VKS, a balance sheet liability for KDI."""
    expense = group_coa_repo.get_by_name(conn, "Interest on  Shareholder loan")
    liability = group_coa_repo.get_by_name(conn, "Loans & Advances taken")
    for code, category in (("VKS", expense), ("KDI", liability)):
        mapping_repo.apply_mapping_for_entity(conn, _entity_id(conn, code), LEDGER,
                                              category["group_account_id"], user="test")
    conn.commit()
    return expense, liability


def _sheet(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Mapping"]
    return wb, ws, {c.value: c.column for c in ws[1]}


def _rows_for(ws, header, ledger):
    return [r for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=header["Ledger Name"]).value == ledger]


def _write_book(tmp_path, name, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    ws.append(header)
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


# --- export ---------------------------------------------------------------------------------

def test_export_gives_a_split_ledger_one_row_per_entity(conn, tmp_path):
    period_id = import_all_entities(conn)
    expense, liability = _split_the_interest_ledger(conn)

    out = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(out))
    _wb, ws, header = _sheet(out)

    rows = _rows_for(ws, header, LEDGER)
    applies = [ws.cell(row=r, column=header["Applies To"]).value for r in rows]
    assert PER_ENTITY in applies, "a split ledger needs a summary row naming it"
    assert "KDI" in applies and "VKS" in applies, "and one editable row per entity"

    by_entity = {ws.cell(row=r, column=header["Applies To"]).value:
                 ws.cell(row=r, column=header["Category"]).value for r in rows}
    assert by_entity["VKS"] == expense["account_name"]
    assert by_entity["KDI"] == liability["account_name"]


def test_export_leaves_an_ordinary_ledger_as_a_single_row(conn, tmp_path):
    """The split layout is for ledgers that need it; everything else stays one line."""
    period_id = import_all_entities(conn)
    out = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(out))
    _wb, ws, header = _sheet(out)

    rows = _rows_for(ws, header, "Sales")
    assert len(rows) == 1
    assert ws.cell(row=rows[0], column=header["Applies To"]).value == ALL_ENTITIES


# --- the regression -------------------------------------------------------------------------

def test_reimporting_an_untouched_export_preserves_the_split(conn, tmp_path):
    """Export then re-import must be a no-op. This is the round trip that used to flatten."""
    period_id = import_all_entities(conn)
    expense, liability = _split_the_interest_ledger(conn)

    out = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(out))
    import_mapping_workbook(conn, str(out), user="test")
    conn.commit()

    assert _category_for(conn, "VKS") == expense["group_account_id"]
    assert _category_for(conn, "KDI") == liability["group_account_id"]


def test_a_whole_ledger_category_on_a_split_ledger_is_refused(conn, tmp_path):
    """An older or hand-edited workbook carrying one row for a split ledger. Refused with an
    explanation rather than applied across every entity."""
    period_id = import_all_entities(conn)
    expense, liability = _split_the_interest_ledger(conn)

    path = _write_book(tmp_path, "hand-edited.xlsx", HEADER,
                       [[LEDGER, ALL_ENTITIES, "KDI, VKS", "", 0, "", liability["account_name"]]])

    with pytest.raises(ValueError, match="booked differently across entities"):
        import_mapping_workbook(conn, str(path), user="test", period_id=period_id)

    assert _category_for(conn, "VKS") == expense["group_account_id"], "nothing may move"
    assert _category_for(conn, "KDI") == liability["group_account_id"]


def test_the_guard_uses_the_database_not_the_file(conn, tmp_path):
    """A ledger can become split after the workbook was exported. The file cannot know that,
    so the split set is recomputed at import time rather than read from the sheet.

    Uses a ledger that starts uniform — the shareholder-loan one is split from the outset by
    its differing Tally primary groups, which is what the export already flags."""
    period_id = import_all_entities(conn)
    uniform = next(
        r for r in mapping_repo.ledger_rows_for_period(conn, period_id)
        if not r["needs_split"] and len({e["entity_code"] for e in r["entities"]}) >= 2
    )
    ledger = uniform["ledger_name"]
    codes = sorted({e["entity_code"] for e in uniform["entities"]})
    sales = group_coa_repo.get_by_name(conn, "Sales")
    bank_charges = group_coa_repo.get_by_name(conn, "Bank Charges")

    # Exported while it was still uniform, so it holds a single ALL_ENTITIES row.
    out = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(out))
    wb, ws, header = _sheet(out)
    rows = _rows_for(ws, header, ledger)
    assert len(rows) == 1
    ws.cell(row=rows[0], column=header["Category"], value=sales["account_name"])
    wb.save(out)

    # ...then someone splits it on the Mapping page before the file is imported back.
    mapping_repo.apply_mapping_for_entity(conn, _entity_id(conn, codes[0]), ledger,
                                          sales["group_account_id"], user="test")
    mapping_repo.apply_mapping_for_entity(conn, _entity_id(conn, codes[1]), ledger,
                                          bank_charges["group_account_id"], user="test")
    conn.commit()

    with pytest.raises(ValueError, match="booked differently across entities"):
        import_mapping_workbook(conn, str(out), user="test", period_id=period_id)
    assert _category_for(conn, codes[1], ledger) == bank_charges["group_account_id"]


def test_a_category_on_the_summary_row_is_refused_rather_than_ignored(conn, tmp_path):
    """Silently dropping it would leave the user believing they had categorized the ledger."""
    period_id = import_all_entities(conn)
    expense, _liability = _split_the_interest_ledger(conn)

    out = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(out))
    wb, ws, header = _sheet(out)
    for r in _rows_for(ws, header, LEDGER):
        if ws.cell(row=r, column=header["Applies To"]).value == PER_ENTITY:
            ws.cell(row=r, column=header["Category"], value=expense["account_name"])
    wb.save(out)

    with pytest.raises(ValueError, match="set the category on its per-entity rows"):
        import_mapping_workbook(conn, str(out), user="test", period_id=period_id)


# --- editing still works ---------------------------------------------------------------------

def test_per_entity_rows_can_still_move_one_entity(conn, tmp_path):
    """The guard blocks the whole-ledger write, not editing. Moving KDI alone must work."""
    period_id = import_all_entities(conn)
    expense, _liability = _split_the_interest_ledger(conn)

    out = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(out))
    wb, ws, header = _sheet(out)
    for r in _rows_for(ws, header, LEDGER):
        if ws.cell(row=r, column=header["Applies To"]).value == "KDI":
            ws.cell(row=r, column=header["Category"], value=expense["account_name"])
    wb.save(out)

    import_mapping_workbook(conn, str(out), user="test", period_id=period_id)
    conn.commit()

    assert _category_for(conn, "KDI") == expense["group_account_id"]
    assert _category_for(conn, "VKS") == expense["group_account_id"]


def test_a_legacy_workbook_without_the_applies_to_column_still_imports(conn, tmp_path):
    """Files exported before the split columns existed keep working — with the guard active,
    so they can update ordinary ledgers but not flatten a split one."""
    period_id = import_all_entities(conn)
    sales = group_coa_repo.get_by_name(conn, "Sales")

    path = _write_book(tmp_path, "legacy.xlsx",
                       ["Ledger Name", "Entities Present In", "Balance (IDR)", "Check", "Category"],
                       [["Sales", "KNS", 0, "", sales["account_name"]]])

    assert import_mapping_workbook(conn, str(path), user="test", period_id=period_id) == 1
    conn.commit()
    assert _category_for(conn, "KNS", "Sales") == sales["group_account_id"]


def test_nothing_is_written_when_any_row_is_rejected(conn, tmp_path):
    """All-or-nothing. The old import applied rows as it went and raised at the end, so an
    invalid category halfway down left the earlier rows already written."""
    period_id = import_all_entities(conn)
    sales = group_coa_repo.get_by_name(conn, "Sales")
    before = _category_for(conn, "KNS", "Sales")

    path = _write_book(tmp_path, "part-bad.xlsx", HEADER, [
        ["Sales", ALL_ENTITIES, "KNS", "", 0, "", sales["account_name"]],
        ["Bank Charges", ALL_ENTITIES, "KNS", "", 0, "", "No Such Category"],
    ])

    with pytest.raises(ValueError, match="unknown category"):
        import_mapping_workbook(conn, str(path), user="test", period_id=period_id)

    assert _category_for(conn, "KNS", "Sales") == before


def test_the_period_is_carried_in_the_file(conn, tmp_path):
    """So the import knows which period's split state to check without being told."""
    period_id = import_all_entities(conn)
    expense, liability = _split_the_interest_ledger(conn)

    path = tmp_path / "Mapping.xlsx"
    export_mapping_workbook(conn, period_id, str(path))

    # No period_id passed: the guard must still fire, from the period stamped at export.
    wb, ws, header = _sheet(path)
    for r in _rows_for(ws, header, LEDGER):
        if ws.cell(row=r, column=header["Applies To"]).value == PER_ENTITY:
            ws.cell(row=r, column=header["Applies To"], value=ALL_ENTITIES)
            ws.cell(row=r, column=header["Category"], value=liability["account_name"])
    wb.save(path)

    with pytest.raises(ValueError, match="booked differently across entities"):
        import_mapping_workbook(conn, str(path), user="test")
    assert _category_for(conn, "VKS") == expense["group_account_id"]
